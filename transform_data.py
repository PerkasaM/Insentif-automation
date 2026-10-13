"""
transform_data.py
Transform data penjualan (semua channel) -> "Raw data by finance <rentang tanggal>.xlsx"

Alur:
  1. Cari file export sales (nama mengandung "export" & "sales", termasuk file retur)
  2. Baca tiap file, BUANG BARIS KUNING (subtotal & grand total), cek total vs grand total
  3. Gabung semua file jadi satu
  4. Transform (dari etl_insentif): clean SDO, drop SO DATE invalid, NET VALUE2,
     CYCLE, BRAND/PG/SUBPG, CHANNEL/SPV/ASM/RSM, REMARKS
  5. Simpan: output\\Raw data by finance 1 April - 30 September.xlsx
     (rentang tanggal diambil dari ACTUAL PGI DATE di dalam data)

Dipanggil oleh run_pipeline.py (via run_insentif.bat), atau langsung:
    python transform_data.py
"""

from __future__ import annotations

import fnmatch
import os
import pickle
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")  # openpyxl berisik soal tanggal kosong (-693595)

# =====================================================================
# 1. CONFIG
# =====================================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUT_DIR = BASE_DIR / "output"
CACHE_DIR = OUT_DIR / ".cache"
OUT_DIR.mkdir(exist_ok=True)

# --- Export sales (sumber transaksi) ---
EXPORT_PATTERN = "*export*sales*.xlsx"          # tidak peka huruf besar/kecil
EXPORT_EXCLUDE = ("gabungan", "zsdconun", "raw data")   # file hasil / sumber lain, jangan dibaca
YELLOW_INDEXED = {43, 13}                       # 43 = subtotal, 13 = grand total
NEGATE_RETUR = True                             # file '*retur*' dijadikan negatif (-> REMARKS RETUR)

# --- Master data ---
FILE_CYCLE = DATA_DIR / "CYCLE_REKAP_1.0.xlsx"
FILE_SKU = DATA_DIR / "skuu6.xlsx"              # sheet MASTERDATA untuk Brand
FILE_CUST = DATA_DIR / "master c09.xlsx"
SHEET_CYCLE = "Sheet2"
SHEET_SKU_MASTER = "MASTERDATA"
SHEET_CUST = "Sheet1"

# --- Aturan bisnis ---
PPN = 1.11
DROP_INVALID_SO_DATE = True        # SO DATE kosong / tidak valid dibuang (sama seperti etl_insentif)
DEDUPE_BILLING = False             # True = 1 baris per DOCNUM + ITEM# + ACTUAL PGI DATE
                                   # (di etl_insentif blok ini di-comment, jadi default False)

# --- Nama file output ---
OUTPUT_PREFIX = "Raw data by finance"
NAMA_FILE_PAKAI_TAHUN = False      # True = selalu tulis tahun (otomatis ditulis jika rentang lintas tahun)

# =====================================================================
# 2. KOLOM
# =====================================================================

HEADER = [
    "TYPE", "CUSTOMER#", "NAME1", "NAME2", "CUST GRP", "SALES GRP", "SDO",
    "DOCNUM", "SO DATE", "DOCDATE", "BILLINGDATE", "ACTUAL PGI DATE",
    "ITEM#", "MATERIAL", "DESCRIPTION", "ITEM CAT", "PLANT", "STORELOC",
    "QUANTITY", "SO AMOUNT", "DISCOUNT", "NET VALUE",
]
DATE_COLS = ["SO DATE", "DOCDATE", "BILLINGDATE", "ACTUAL PGI DATE"]
NUM_COLS = ["QUANTITY", "SO AMOUNT", "DISCOUNT", "NET VALUE"]
TEXT_COLS = [c for c in HEADER if c not in DATE_COLS + NUM_COLS]

OUTPUT_COLS = [
    "TYPE", "CUSTOMER#", "NAME1", "NAME2", "CUST GRP", "SALES GRP", "SDO",
    "DOCNUM", "SO DATE", "DOCDATE", "BILLINGDATE", "ACTUAL PGI DATE",
    "ITEM#", "MATERIAL", "DESCRIPTION", "ITEM CAT", "PLANT", "STORELOC",
    "QUANTITY", "SO AMOUNT", "DISCOUNT", "NET VALUE", "NET VALUE2",
    "CYCLE", "BRAND", "PG", "SUBPG", "SUBPG1", "CHANNEL", "REMARKS",
    "CUST_TYPE", "SPV", "ASM", "RSM",
]

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]


# =====================================================================
# 3. BACA & GABUNG EXPORT SALES (buang baris kuning)
# =====================================================================

def _is_yellow(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    color = fill.fgColor
    return color.type == "indexed" and color.indexed in YELLOW_INDEXED


def _to_number(x):
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).replace(",", "").strip())
    except ValueError:
        return None


def _to_date(x):
    """Tanggal valid -> datetime; error / kosong / 0000-00-00 -> None."""
    if isinstance(x, datetime):
        return x if x.year > 1900 else None
    return None


def _find_export_files() -> list[str]:
    """Cari file export sales di folder data, kalau tidak ada coba folder script."""
    for folder in (DATA_DIR, BASE_DIR):
        if not folder.exists():
            continue
        files = sorted(
            str(folder / n) for n in os.listdir(folder)
            if fnmatch.fnmatch(n.lower(), EXPORT_PATTERN.lower())
            and not n.startswith("~$")
            and not any(p in n.lower() for p in EXPORT_EXCLUDE)
        )
        if files:
            return files
    raise FileNotFoundError(
        f"Tidak ada file export sales ('{EXPORT_PATTERN}') di {DATA_DIR} maupun {BASE_DIR}"
    )


def _read_export_file(path: str) -> pd.DataFrame:
    """Baca 1 file export, buang baris kuning, cek total detail vs grand total."""
    ws = openpyxl.load_workbook(path).active
    detail, grand = [], None
    for row in ws.iter_rows(min_row=1):
        values = [c.value for c in row]
        if values[0] == "TYPE":                          # baris header
            continue
        if _is_yellow(row[0]):
            if values[1] in ("", None) and values[7] in ("", None):
                grand = values[18:22]                    # grand total
            continue                                     # subtotal / grand total -> buang
        detail.append(values)

    df = pd.DataFrame(detail, columns=HEADER)
    for c in TEXT_COLS:
        df[c] = df[c].apply(lambda v: "" if v is None else re.sub(r"\s+", " ", str(v)).strip())
    for c in DATE_COLS:
        df[c] = pd.to_datetime(df[c].apply(_to_date))
    for c in NUM_COLS:
        df[c] = df[c].apply(_to_number)

    if grand:
        g = [_to_number(v) for v in grand]
        s = [df[c].sum() for c in NUM_COLS]
        if not all(abs((a or 0) - b) < 1 for a, b in zip(g, s)):
            raise ValueError(
                f"{os.path.basename(path)}: total detail tidak cocok dengan grand total "
                f"(grand={g}, detail={s}) - cek deteksi baris kuning"
            )
    return df


def load_export_sales() -> pd.DataFrame:
    """Gabung semua file export sales + retur menjadi satu DataFrame."""
    files = _find_export_files()
    CACHE_DIR.mkdir(parents=True, exist_ok=True)

    frames = []
    for f in files:
        name = os.path.basename(f)
        st = os.stat(f)
        cache_file = CACHE_DIR / f"{name}.{st.st_size}.{int(st.st_mtime)}.pkl"

        if cache_file.exists():
            with open(cache_file, "rb") as fh:
                df = pickle.load(fh)
            src = "cache"
        else:
            df = _read_export_file(f)
            for old in CACHE_DIR.glob(f"{name}.*.pkl"):
                old.unlink()
            with open(cache_file, "wb") as fh:
                pickle.dump(df, fh)
            src = "baca"

        is_retur = "retur" in name.lower()
        if is_retur and NEGATE_RETUR:
            for c in NUM_COLS:
                df[c] = -df[c].abs()
        note = ", retur dinegatifkan" if (is_retur and NEGATE_RETUR) else ""
        print(f"   - {name:<45s} {len(df):>8,} baris ({src}{note})")
        frames.append(df)

    data = pd.concat(frames, ignore_index=True)
    print(f"   digabung: {len(data):,} baris dari {len(files)} file")
    return data


# =====================================================================
# 4. TRANSFORM (dari etl_insentif)
# =====================================================================

def cleaning_sdo_global(s):
    """Normalisasi nama SDO."""
    if pd.isna(s):
        return None
    s = str(s).upper().strip()
    if not s:
        return None
    if "ANDY WIJAYA" in s:
        return "ANDY"
    elif "IMAM SHOVI" in s:
        return "IMAM SHOVII"
    elif "OKTO" in s:
        return "OKTO RISMAN"
    elif "ANDI MANGGALA" in s or "ANDI ARDIYANSYAH" in s:
        return "ANDI MANGGALA PUTRA"
    elif "JINO" in s:
        return "FREEKY JINO"
    elif "MUHAMMAD LANANG GALIH" in s:
        return "MUHAMMAD LANANG GALIH GUMILANG"
    elif "ABDUL JABBAR" in s:
        return "ABDUL JABAR"
    elif "MO MOHAMMAD HENDRA FARIZAL" in s:
        return "MOHAMMAD HENDRA FARIZAL"
    else:
        return s


def _clean_transaksi(df: pd.DataFrame) -> pd.DataFrame:
    """Drop baris invalid + normalisasi tipe data + cleaning SDO."""
    df = df.copy()

    # 1) Cleaning SDO di awal
    df["SDO"] = df["SDO"].apply(cleaning_sdo_global)

    # 2) Drop baris dengan SO DATE kosong / tidak valid
    if DROP_INVALID_SO_DATE:
        mask_invalid = df["SO DATE"].isna()
        n_drop = int(mask_invalid.sum())
        if n_drop:
            nilai = df.loc[mask_invalid, "NET VALUE"].sum()
            print(f"   drop {n_drop:,} baris dengan SO DATE invalid (net value {nilai:,.0f})")
            df = df[~mask_invalid].copy()

    # 3) ACTUAL PGI DATE -> datetime
    df["ACTUAL PGI DATE"] = pd.to_datetime(df["ACTUAL PGI DATE"], errors="coerce")
    df = df.dropna(subset=["ACTUAL PGI DATE"]).reset_index(drop=True)

    # 4) Normalisasi key
    df["CUSTOMER#"] = df["CUSTOMER#"].astype(str).str.strip()
    df["MATERIAL"] = df["MATERIAL"].astype(str).str.strip()
    return df


def _build_cycle_map(cycle_df: pd.DataFrame) -> pd.DataFrame:
    c = cycle_df[["Cycle", "Start Date", "End Date"]].dropna(subset=["Cycle"]).copy()
    c["Start Date"] = pd.to_datetime(c["Start Date"], errors="coerce")
    c["End Date"] = pd.to_datetime(c["End Date"], errors="coerce")
    return c.dropna(subset=["Start Date", "End Date"]).reset_index(drop=True)


def _assign_cycle(df: pd.DataFrame, cycle_df: pd.DataFrame) -> pd.Series:
    """Cari cycle yang Start-End-nya memuat ACTUAL PGI DATE."""
    c_sorted = cycle_df.sort_values("Start Date").reset_index(drop=True)
    starts = c_sorted["Start Date"].values
    ends = c_sorted["End Date"].values
    cycles = c_sorted["Cycle"].values
    pgi = df["ACTUAL PGI DATE"].values

    idx = np.searchsorted(starts, pgi, side="right") - 1
    result = []
    for i, ix in enumerate(idx):
        if 0 <= ix < len(cycles) and starts[ix] <= pgi[i] <= ends[ix]:
            result.append(cycles[ix])
        else:
            result.append(None)
    return pd.Series(result, index=df.index, name="CYCLE")


def _prepare_sku_master(sku_df: pd.DataFrame) -> pd.DataFrame:
    s = sku_df[["SKU Code", "Brand", "PG", "SUB-PG"]].copy()
    s.columns = ["SKU Code", "BRAND_SRC", "PG_SRC", "SUBPG_SRC"]
    s["SKU Code"] = s["SKU Code"].astype(str).str.strip()
    s = s.dropna(subset=["SKU Code"])
    return s.drop_duplicates(subset=["SKU Code"], keep="first").reset_index(drop=True)


def _prepare_cust(cust_df: pd.DataFrame) -> pd.DataFrame:
    c = cust_df[["Customer Code", "Chanel", "SPV", "ASM", "RSM"]].copy()
    c.columns = ["CUSTOMER#", "CHANNEL_SRC", "SPV_SRC", "ASM_SRC", "RSM_SRC"]
    c["CUSTOMER#"] = c["CUSTOMER#"].astype(str).str.strip()
    c = c.dropna(subset=["CUSTOMER#"])
    return c.drop_duplicates(subset=["CUSTOMER#"], keep="first").reset_index(drop=True)


def _classify_remarks(netvalue2: float):
    """<0 RETUR, <1000 CLAIM G, >=1000 SALES."""
    if pd.isna(netvalue2):
        return None
    if netvalue2 < 0:
        return "RETUR"
    if netvalue2 < 1000:
        return "CLAIM G"
    return "SALES"


def load_masters() -> dict:
    print("   Membaca master data (cycle, SKU, customer)...")
    for f in (FILE_CYCLE, FILE_SKU, FILE_CUST):
        if not f.exists():
            raise FileNotFoundError(f"File master tidak ditemukan: {f}")
    return {
        "cycle": pd.read_excel(FILE_CYCLE, sheet_name=SHEET_CYCLE, engine="openpyxl"),
        "sku_master": pd.read_excel(FILE_SKU, sheet_name=SHEET_SKU_MASTER, engine="openpyxl"),
        "cust": pd.read_excel(FILE_CUST, sheet_name=SHEET_CUST, engine="openpyxl"),
    }


def transform(raw: pd.DataFrame, masters: dict) -> pd.DataFrame:
    """Transaksi export sales -> data enriched."""
    df = _clean_transaksi(raw)
    print(f"   baris setelah cleaning: {len(df):,}")

    # 1) NET VALUE2 = NET VALUE x PPN
    df["NET VALUE2"] = df["NET VALUE"] * PPN

    # 2) CYCLE dari ACTUAL PGI DATE
    cycle_map = _build_cycle_map(masters["cycle"])
    df["CYCLE"] = _assign_cycle(df, cycle_map)
    n_ok, n_all = int(df["CYCLE"].notna().sum()), len(df)
    print(f"   cycle terisi: {n_ok:,} / {n_all:,} ({n_ok / n_all * 100:.1f}%)")

    # 3) BRAND, PG, SUBPG dari MATERIAL
    sku = _prepare_sku_master(masters["sku_master"])
    df = df.merge(sku, how="left", left_on="MATERIAL", right_on="SKU Code")
    df["BRAND"] = df["BRAND_SRC"]
    df["PG"] = df["PG_SRC"]
    df["SUBPG"] = df["SUBPG_SRC"]
    df["SUBPG1"] = df["SUBPG_SRC"]
    df = df.drop(columns=["SKU Code", "BRAND_SRC", "PG_SRC", "SUBPG_SRC"], errors="ignore")

    # 4) CHANNEL, SPV, ASM, RSM dari CUSTOMER#
    cust = _prepare_cust(masters["cust"])
    df = df.merge(cust, how="left", on="CUSTOMER#")
    df["CHANNEL"] = df["CHANNEL_SRC"]
    df["SPV"] = df["SPV_SRC"]
    df["ASM"] = df["ASM_SRC"]
    df["RSM"] = df["RSM_SRC"]
    df = df.drop(columns=["CHANNEL_SRC", "SPV_SRC", "ASM_SRC", "RSM_SRC"], errors="ignore")

    # 5) REMARKS & CUST_TYPE
    df["REMARKS"] = df["NET VALUE2"].apply(_classify_remarks)
    df["CUST_TYPE"] = None

    df = df.reindex(columns=OUTPUT_COLS)

    # 6) (opsional) dedupe billing triplicates
    if DEDUPE_BILLING:
        before = len(df)
        df = df.drop_duplicates(subset=["DOCNUM", "ITEM#", "ACTUAL PGI DATE"], keep="first")
        print(f"   dedupe billing: -{before - len(df):,} baris")

    # Info kelengkapan master
    print(f"   BRAND kosong  : {int(df['BRAND'].isna().sum()):,} baris (MATERIAL tidak ada di master SKU)")
    print(f"   CHANNEL kosong: {int(df['CHANNEL'].isna().sum()):,} baris (CUSTOMER# tidak ada di master customer)")
    return df


# =====================================================================
# 5. SIMPAN "Raw data by finance <rentang>"
# =====================================================================

def make_range_label(df: pd.DataFrame) -> str:
    """Contoh: '1 April - 30 September' (dari ACTUAL PGI DATE)."""
    d1, d2 = df["ACTUAL PGI DATE"].min(), df["ACTUAL PGI DATE"].max()
    pakai_tahun = NAMA_FILE_PAKAI_TAHUN or d1.year != d2.year
    t1 = f"{d1.day} {BULAN[d1.month - 1]}" + (f" {d1.year}" if pakai_tahun else "")
    t2 = f"{d2.day} {BULAN[d2.month - 1]}" + (f" {d2.year}" if pakai_tahun else "")
    return f"{t1} - {t2}"


def _excel_engine() -> str:
    """xlsxwriter jauh lebih cepat untuk data besar; openpyxl sebagai cadangan."""
    try:
        import xlsxwriter  # noqa: F401
        return "xlsxwriter"
    except ImportError:
        return "openpyxl"


def save_raw(df: pd.DataFrame, label: str) -> Path:
    path = OUT_DIR / f"{OUTPUT_PREFIX} {label}.xlsx"
    engine = _excel_engine()
    print(f"   Menyimpan {path.name} ({len(df):,} baris)...")
    if engine == "openpyxl":
        print("   (lambat: jalankan 'pip install xlsxwriter' agar penyimpanan jauh lebih cepat)")
    try:
        with pd.ExcelWriter(path, engine=engine,
                            datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as w:
            df.to_excel(w, sheet_name="Raw Data", index=False)
            ws = w.sheets["Raw Data"]
            if engine == "xlsxwriter":
                ws.freeze_panes(1, 0)
                ws.autofilter(0, 0, len(df), len(df.columns) - 1)
                for i, col in enumerate(df.columns):
                    ws.set_column(i, i, min(max(len(col) + 3, 11), 30))
            else:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for i, col in enumerate(df.columns, start=1):
                    ws.cell(row=1, column=i).font = Font(bold=True)
                    ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 3, 11), 30)
    except PermissionError:
        raise PermissionError(f"Tidak bisa menulis {path.name} - tutup dulu file itu di Excel lalu jalankan ulang.")
    return path


def find_latest_raw() -> Path | None:
    files = sorted(OUT_DIR.glob(f"{OUTPUT_PREFIX}*.xlsx"), key=lambda p: p.stat().st_mtime)
    return files[-1] if files else None


def load_latest_raw() -> tuple[pd.DataFrame, str]:
    """Untuk menjalankan script channel sendiri-sendiri tanpa transform ulang."""
    p = find_latest_raw()
    if p is None:
        raise FileNotFoundError(f"Belum ada file '{OUTPUT_PREFIX} ...' di {OUT_DIR}. Jalankan transform dulu.")
    print(f"   Membaca {p.name} ...")
    df = pd.read_excel(p, sheet_name="Raw Data", engine="openpyxl")
    return df, p.stem.replace(OUTPUT_PREFIX, "", 1).strip()


# =====================================================================
# 6. RUN
# =====================================================================

def run() -> tuple[pd.DataFrame, Path, str]:
    """Jalankan seluruh transform. Return (df, path_file_raw, label_rentang)."""
    print("[Transform 1/3] Membaca & menggabung export sales (baris kuning dibuang)...")
    raw = load_export_sales()

    print("[Transform 2/3] Transform data...")
    masters = load_masters()
    df = transform(raw, masters)
    if df.empty:
        raise ValueError("Hasil transform kosong - cek file export sales dan filter SO DATE.")

    print("[Transform 3/3] Menyimpan hasil transform...")
    label = make_range_label(df)
    path = save_raw(df, label)
    print(f"   Tersimpan: {path}")
    return df, path, label


if __name__ == "__main__":
    try:
        run()
    except Exception as e:  # pesan singkat untuk user .bat
        print(f"\n[GAGAL] {e}")
        sys.exit(1)
